import pytest

from osintlab import http
from osintlab.config import get_settings
from osintlab.ingest import documents
from osintlab.plugins import load_all


def test_bloqueio_offline_e_ssrf():
    with pytest.raises(http.FetchBlocked):
        http.check_url_allowed("https://exemplo.com")
    get_settings().offline = False
    with pytest.raises(http.FetchBlocked):
        http.check_url_allowed("http://127.0.0.1:8080/admin")
    with pytest.raises(http.FetchBlocked):
        http.check_url_allowed("file:///etc/passwd")


def test_conectores_registrados_e_passivos():
    reg = load_all()
    for nome in ("dns", "rdap", "crtsh", "wayback", "cnpj", "gdelt", "geocodificar", "rss", "querido_diario"):
        assert nome in reg and reg[nome].passive


def test_extracao_docx_xlsx_pdf_json():
    import io
    import docx
    import openpyxl
    d = docx.Document(); d.add_paragraph("Relatório da Empresa Beta Ltda"); b = io.BytesIO(); d.save(b)
    assert "Beta" in documents.extract(b.getvalue(), "a.docx").text
    wb = openpyxl.Workbook(); wb.active.append(["col", 1]); b = io.BytesIO(); wb.save(b)
    ex = documents.extract(b.getvalue(), "a.xlsx")
    assert ex.kind == "dataset" and "col" in ex.text
    from reportlab.pdfgen import canvas
    b = io.BytesIO(); c = canvas.Canvas(b); c.drawString(100, 750, "Documento PDF sintetico"); c.save()
    assert "sintetico" in documents.extract(b.getvalue(), "a.pdf").text
    assert "chave: valor" in documents.extract(b'{"chave": "valor"}', "a.json").text


def test_html_metadados():
    from osintlab.ingest.web import extract_html
    ex = extract_html('<html lang="pt"><head><title>T</title><meta property="article:published_time" content="2024-01-02">'
                      '</head><body><a href="https://x.example/a">x</a><script>alert(1)</script><p>Texto</p></body></html>')
    assert ex.title == "T" and ex.metadata["publicado_em"] == "2024-01-02" and "alert" not in ex.text
