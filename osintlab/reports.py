"""Relatórios: executivo, técnico, cronológico, geográfico, de monitoramento e manifesto.

Formatos: HTML (autocontido), JSON, Markdown, DOCX e PDF.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, select_autoescape

from . import audit, claims, evidence, monitor
from .analysis import compare, correlation, geo, timeline, verification
from .config import get_settings
from .core import get_project, list_entities, list_items, stats
from .db import now_iso

TIPOS = {"executivo": "Relatório executivo", "tecnico": "Relatório técnico", "cronologico": "Relatório cronológico",
         "geografico": "Relatório geográfico", "monitoramento": "Relatório de monitoramento",
         "completo": "Relatório completo"}

AVISO = ("Documento produzido em projeto pessoal, experimental e não comercial de pesquisa em fontes abertas. "
         "As correspondências entre entidades distinguem vínculo confirmado, possível relação, coincidência e hipótese; "
         "nenhuma hipótese deve ser tratada como conclusão. Verifique as fontes antes de qualquer uso.")


def collect(project_id: int, tipo: str = "completo") -> dict:
    p = get_project(project_id)
    if not p:
        raise KeyError("projeto inexistente")
    sec = {
        "executivo": {"resumo", "afirmacoes", "itens", "relacoes_fortes"},
        "tecnico": {"resumo", "afirmacoes", "itens", "entidades", "relacoes", "verificacao", "evidencias"},
        "cronologico": {"resumo", "linha_do_tempo", "afirmacoes"},
        "geografico": {"resumo", "geo"},
        "monitoramento": {"resumo", "monitoramento"},
    }.get(tipo) or {"resumo", "afirmacoes", "itens", "entidades", "relacoes", "verificacao", "linha_do_tempo", "geo",
                    "monitoramento", "evidencias"}
    data = {"tipo": tipo, "titulo": TIPOS.get(tipo, tipo), "projeto": p, "gerado_em": now_iso(), "aviso": AVISO,
            "estatisticas": stats(project_id), "secoes": sorted(sec)}
    if "afirmacoes" in sec:
        cl = claims.list_(project_id)
        data["afirmacoes"] = {k: [c for c in cl if c["kind"] == k] for k in claims.KINDS}
        data["naturezas"] = claims.KINDS
    if tipo in ("executivo", "completo"):
        from .ai.assistant import resumo_guardado
        data["resumo_ia"] = resumo_guardado(project_id)
    if "itens" in sec:
        data["itens"] = list_items(project_id, limit=60 if tipo != "executivo" else 15)
    if "entidades" in sec:
        data["entidades"] = [e for e in list_entities(project_id, include_sensitive=False)][:300]
    if "relacoes" in sec or "relacoes_fortes" in sec:
        from .db import get_db
        sens = {e["id"] for e in get_db().query("SELECT id FROM entities WHERE project_id=? AND sensitive=1", (project_id,))}
        rels = [r for r in correlation.relations(project_id) if r["src_entity"] not in sens and r["dst_entity"] not in sens]
        data["relacoes"] = [r for r in rels if "relacoes" in sec or r["level"] in ("confirmado", "possivel")][:400]
        data["niveis"] = correlation.LEVELS
    if "verificacao" in sec:
        data["verificacao"] = {"temporais": verification.temporal_checks(project_id),
                               "circulares": verification.circular_citations(project_id),
                               "duplicatas": compare.near_duplicates(project_id)[:50]}
    if "linha_do_tempo" in sec:
        data["linha_do_tempo"] = [e for e in timeline.build(project_id) if e["tipo"] != "coleta"][:1000]
    if "geo" in sec:
        fc = geo.features(project_id)
        data["geo"] = {"pontos": [{"lat": f["geometry"]["coordinates"][1], "lon": f["geometry"]["coordinates"][0],
                                   **f["properties"]} for f in fc["features"]], "distancias": geo.distance_matrix(fc)[:30]}
    if "monitoramento" in sec:
        data["monitoramento"] = monitor.events(project_id, limit=300)
    if "evidencias" in sec:
        data["manifesto"] = evidence.manifest(project_id)
    data["metodologia"] = [
        "Coleta exclusivamente passiva de fontes abertas; originais preservados com SHA-256 e cabeçalhos HTTP.",
        "Confiabilidade de fontes segundo o Código do Almirantado (A–F / 1–6).",
        "Corroboração considera apenas fontes independentes (domínios distintos, sem cópia de texto ou citação mútua).",
        "Relações automáticas: confirmadas apenas por regras determinísticas; coocorrências viram 'possível relação' "
        "ou 'coincidência'; semelhança de nomes vira 'hipótese'.",
        "Respostas de IA restritas a trechos recuperados; citações inexistentes são removidas automaticamente.",
    ]
    return data


HTML_TMPL = """<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<title>{{ d.titulo }} — {{ d.projeto.name }}</title>
<style>
body{font-family:system-ui,Segoe UI,Arial,sans-serif;max-width:1100px;margin:2rem auto;padding:0 1rem;color:#1c2330;line-height:1.5}
h1{font-size:1.6rem;margin-bottom:.2rem}h2{border-bottom:2px solid #d7dde8;padding-bottom:.3rem;margin-top:2rem}
.meta{color:#5b6576;font-size:.9rem}.aviso{background:#fff7e0;border-left:4px solid #e0a800;padding:.6rem 1rem;font-size:.9rem}
table{border-collapse:collapse;width:100%;font-size:.85rem;margin:.5rem 0}th,td{border:1px solid #d7dde8;padding:.3rem .5rem;text-align:left;vertical-align:top}
th{background:#f0f3f8}.tag{display:inline-block;padding:0 .4rem;border-radius:4px;font-size:.75rem;background:#e8edf5}
.confirmado{background:#d4f4dd}.possivel{background:#dfe9ff}.coincidencia{background:#eee}.hipotese{background:#ffe9c7}
code{font-size:.8rem;word-break:break-all}.kpi{display:inline-block;margin:.3rem 1rem .3rem 0}
</style></head><body>
<h1>{{ d.titulo }}: {{ d.projeto.name }}</h1>
<div class="meta">Gerado em {{ d.gerado_em }} · Finalidade declarada: {{ d.projeto.purpose }} · Base: {{ d.projeto.legal_basis }}</div>
<p class="aviso">{{ d.aviso }}</p>
<h2>Resumo</h2>
{% for k, v in d.estatisticas.items() %}<span class="kpi"><b>{{ v }}</b> {{ k }}</span>{% endfor %}
{% if d.projeto.description %}<p>{{ d.projeto.description }}</p>{% endif %}
{% if d.resumo_ia %}<h2>Resumo executivo</h2><p class="meta">Redigido por IA ({{ d.resumo_ia.modo }}) em {{ d.resumo_ia.gerado_em }}
a partir das fontes do caso; os números entre colchetes remetem às fontes listadas abaixo. Revise antes de usar.</p>
<div style="white-space:pre-wrap">{{ d.resumo_ia.texto }}</div>
<ol style="font-size:.85rem">{% for f in d.resumo_ia.fontes %}<li>{% if f.url %}<a href="{{ f.url }}">{{ f.titulo or f.url }}</a>{% else %}{{ f.titulo }}{% endif %} (#{{ f.item }}, {{ f.fonte }}{% if f.data %}, {{ f.data[:10] }}{% endif %})</li>{% endfor %}</ol>{% endif %}
{% if d.afirmacoes is defined %}<h2>Afirmações e hipóteses</h2>
{% for k, lst in d.afirmacoes.items() if lst %}<h3>{{ d.naturezas[k] }}</h3><ul>
{% for c in lst %}<li>{{ c.statement }} {% if c.supporting %}<span class="tag">apoio: {% for s in c.supporting %}#{{ s }} {% endfor %}</span>{% endif %}
{% if c.contradicting %}<span class="tag hipotese">contra: {% for s in c.contradicting %}#{{ s }} {% endfor %}</span>{% endif %}
{% if c.confidence %}<span class="tag">{{ c.confidence }}</span>{% endif %}{% if c.rationale %}<br><small>{{ c.rationale }}</small>{% endif %}</li>{% endfor %}</ul>
{% else %}<p>Nenhuma afirmação registrada.</p>{% endfor %}{% endif %}
{% if d.itens is defined %}<h2>Fontes e itens (por triagem)</h2><table><tr><th>#</th><th>Título</th><th>Fonte</th><th>Conf.</th><th>Data</th><th>Pont.</th><th>Evidência</th></tr>
{% for i in d.itens %}<tr><td>{{ i.id }}</td><td>{% if i.url %}<a href="{{ i.url }}">{{ i.title or i.url }}</a>{% else %}{{ i.title }}{% endif %}</td>
<td>{{ i.source_domain or 'local' }}</td><td>{{ i.source_reliability }}{{ i.info_credibility }}</td><td>{{ (i.published_at or '')[:10] }}</td><td>{{ i.triage_score }}</td>
<td><code>{{ (i.evidence_sha256 or '')[:16] }}</code></td></tr>{% endfor %}</table>{% endif %}
{% if d.relacoes is defined %}<h2>Vínculos</h2><table><tr><th>Nível</th><th>Origem</th><th>Relação</th><th>Destino</th><th>Justificativa</th></tr>
{% for r in d.relacoes %}<tr><td><span class="tag {{ r.level }}">{{ d.niveis[r.level] }}</span></td><td>{{ r.src_type }}: {{ r.src_value }}</td>
<td>{{ r.rel_type }}</td><td>{{ r.dst_type }}: {{ r.dst_value }}</td><td>{{ r.rationale }}</td></tr>{% endfor %}</table>{% endif %}
{% if d.entidades is defined %}<h2>Entidades (não sensíveis)</h2><table><tr><th>Tipo</th><th>Valor</th><th>Itens</th></tr>
{% for e in d.entidades %}<tr><td>{{ e.type }}</td><td>{{ e.value }}</td><td>{{ e.n_itens }}</td></tr>{% endfor %}</table>{% endif %}
{% if d.verificacao is defined %}<h2>Verificação</h2>
<h3>Inconsistências temporais</h3>{% if d.verificacao.temporais %}<ul>{% for t in d.verificacao.temporais %}<li>#{{ t.item }} {{ t.titulo }} — {{ t.problema }}</li>{% endfor %}</ul>{% else %}<p>Nenhuma detectada.</p>{% endif %}
<h3>Citações circulares</h3>{% if d.verificacao.circulares %}<ul>{% for c in d.verificacao.circulares %}<li>{{ c|join(' → ') }} → {{ c[0] }}</li>{% endfor %}</ul>{% else %}<p>Nenhuma detectada.</p>{% endif %}
<h3>Conteúdo quase duplicado</h3>{% if d.verificacao.duplicatas %}<ul>{% for x in d.verificacao.duplicatas %}<li>#{{ x.a }} ({{ x.a_fonte }}) ≈ #{{ x.b }} ({{ x.b_fonte }}): {{ (x.similaridade*100)|round(0) }}%</li>{% endfor %}</ul>{% else %}<p>Nenhum.</p>{% endif %}{% endif %}
{% if d.linha_do_tempo is defined %}<h2>Linha do tempo</h2><table><tr><th>Data</th><th>Tipo</th><th>Descrição</th><th>Natureza</th><th>Item</th></tr>
{% for e in d.linha_do_tempo %}<tr><td>{{ e.data }}</td><td>{{ e.tipo }}</td><td>{{ e.descricao }}</td><td>{{ e.natureza }}</td><td>{{ e.item or '' }}</td></tr>{% endfor %}</table>{% endif %}
{% if d.geo is defined %}<h2>Dados geográficos</h2><table><tr><th>Lat</th><th>Lon</th><th>Origem da coordenada</th><th>Item</th><th>Rótulo</th></tr>
{% for g in d.geo.pontos %}<tr><td>{{ g.lat }}</td><td>{{ g.lon }}</td><td>{{ g.fonte_coordenada }}</td><td>#{{ g.item }} {{ g.titulo }}</td><td>{{ g.rotulo }}</td></tr>{% endfor %}</table>{% endif %}
{% if d.monitoramento is defined %}<h2>Monitoramento</h2><table><tr><th>Data</th><th>Alvo</th><th>Tipo</th><th>Resumo</th></tr>
{% for e in d.monitoramento %}<tr><td>{{ e.ts }}</td><td>{{ e.target }}</td><td>{{ e.change_type }}</td><td>{{ e.summary }}</td></tr>{% endfor %}</table>{% endif %}
{% if d.manifesto is defined %}<h2>Manifesto de evidências</h2><p>Total: {{ d.manifesto.total }} · Hash raiz (SHA-256): <code>{{ d.manifesto.hash_raiz }}</code> ·
Trilha de auditoria íntegra: <b>{{ 'sim' if d.manifesto.auditoria_integra else 'NÃO' }}</b></p>
<table><tr><th>SHA-256</th><th>Origem</th><th>Coletado em</th><th>Íntegro</th></tr>{% for e in d.manifesto.evidencias %}
<tr><td><code>{{ e.sha256 }}</code></td><td>{{ e.url_origem or e.nome_original }}</td><td>{{ e.coletado_em }}</td><td>{{ 'sim' if e.integro else 'NÃO' }}</td></tr>{% endfor %}</table>{% endif %}
<h2>Metodologia e limitações</h2><ul>{% for m in d.metodologia %}<li>{{ m }}</li>{% endfor %}</ul>
</body></html>"""


def render_html(data: dict) -> str:
    env = Environment(autoescape=select_autoescape(["html"]))
    return env.from_string(HTML_TMPL).render(d=data)


def render_markdown(data: dict) -> str:
    p = data["projeto"]
    out = [f"# {data['titulo']}: {p['name']}", "", f"_Gerado em {data['gerado_em']}. Finalidade: {p['purpose']}_", "",
           f"> {data['aviso']}", "", "## Resumo", ""]
    out += [f"- **{k}**: {v}" for k, v in data["estatisticas"].items()]
    if data.get("resumo_ia"):
        r = data["resumo_ia"]
        out += ["", "## Resumo executivo", "", f"_Redigido por IA ({r['modo']}) em {r['gerado_em']}; revise antes de usar._",
                "", *[ln.replace("## ", "### ", 1) if ln.startswith("## ") else ln for ln in r["texto"].splitlines()],
                "", "Fontes do resumo:", ""]
        out += [f"{f['n']}. {f['titulo'] or f['url'] or ''} (#{f['item']}, {f['fonte']})" for f in r["fontes"]]
    if "afirmacoes" in data:
        out += ["", "## Afirmações e hipóteses"]
        for k, lst in data["afirmacoes"].items():
            if lst:
                out += ["", f"### {data['naturezas'][k]}"] + [
                    f"- {c['statement']} (apoio: {', '.join('#' + str(s) for s in c['supporting']) or '—'})" for c in lst]
    if "itens" in data:
        out += ["", "## Itens", "", "| # | Título | Fonte | Conf. | Pont. |", "|---|---|---|---|---|"]
        out += [f"| {i['id']} | {(i['title'] or '').replace('|', '/')[:80]} | {i['source_domain'] or 'local'} | "
                f"{i['source_reliability']}{i['info_credibility']} | {i['triage_score']} |" for i in data["itens"]]
    if "relacoes" in data:
        out += ["", "## Vínculos", ""] + [f"- [{data['niveis'][r['level']]}] {r['src_value']} —{r['rel_type']}→ "
                                          f"{r['dst_value']}: {r['rationale']}" for r in data["relacoes"]]
    if "linha_do_tempo" in data:
        out += ["", "## Linha do tempo", ""] + [f"- {e['data']} ({e['tipo']}): {e['descricao'][:200]}"
                                                for e in data["linha_do_tempo"]]
    if "manifesto" in data:
        out += ["", "## Manifesto de evidências", "", f"Hash raiz: `{data['manifesto']['hash_raiz']}`", ""]
        out += [f"- `{e['sha256']}` {e['url_origem'] or e['nome_original']}" for e in data["manifesto"]["evidencias"]]
    out += ["", "## Metodologia", ""] + [f"- {m}" for m in data["metodologia"]]
    return "\n".join(out)


def render_docx(data: dict, path: Path) -> None:
    import docx
    d = docx.Document()
    p = data["projeto"]
    d.add_heading(f"{data['titulo']}: {p['name']}", 0)
    d.add_paragraph(f"Gerado em {data['gerado_em']} · Finalidade: {p['purpose']}")
    d.add_paragraph(data["aviso"]).italic = True
    d.add_heading("Resumo", 1)
    for k, v in data["estatisticas"].items():
        d.add_paragraph(f"{k}: {v}", style="List Bullet")
    if data.get("resumo_ia"):
        r = data["resumo_ia"]
        d.add_heading("Resumo executivo", 1)
        d.add_paragraph(f"Redigido por IA ({r['modo']}) em {r['gerado_em']}; revise antes de usar.").italic = True
        for ln in r["texto"].splitlines():
            if ln.startswith("## "):
                d.add_heading(ln[3:], 2)
            elif ln.strip():
                d.add_paragraph(ln.lstrip("- "))
        for f in r["fontes"]:
            d.add_paragraph(f"[{f['n']}] {f['titulo'] or f['url'] or ''} (#{f['item']}, {f['fonte']})", style="List Number")
    if "afirmacoes" in data:
        d.add_heading("Afirmações e hipóteses", 1)
        for k, lst in data["afirmacoes"].items():
            if lst:
                d.add_heading(data["naturezas"][k], 2)
                for c in lst:
                    d.add_paragraph(f"{c['statement']} (apoio: {', '.join('#' + str(s) for s in c['supporting']) or '—'})",
                                    style="List Bullet")
    if "itens" in data:
        d.add_heading("Itens", 1)
        t = d.add_table(rows=1, cols=5)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ("#", "Título", "Fonte", "Conf.", "Pont.")):
            c.text = h
        for i in data["itens"]:
            row = t.add_row().cells
            row[0].text, row[1].text = str(i["id"]), (i["title"] or "")[:120]
            row[2].text, row[3].text = i["source_domain"] or "local", f"{i['source_reliability']}{i['info_credibility']}"
            row[4].text = str(i["triage_score"])
    if "relacoes" in data:
        d.add_heading("Vínculos", 1)
        for r in data["relacoes"][:200]:
            d.add_paragraph(f"[{data['niveis'][r['level']]}] {r['src_value']} —{r['rel_type']}→ {r['dst_value']}. "
                            f"{r['rationale']}", style="List Bullet")
    if "linha_do_tempo" in data:
        d.add_heading("Linha do tempo", 1)
        for e in data["linha_do_tempo"][:300]:
            d.add_paragraph(f"{e['data']} ({e['tipo']}): {e['descricao'][:200]}", style="List Bullet")
    if "manifesto" in data:
        d.add_heading("Manifesto de evidências", 1)
        d.add_paragraph(f"Hash raiz SHA-256: {data['manifesto']['hash_raiz']}")
        for e in data["manifesto"]["evidencias"]:
            d.add_paragraph(f"{e['sha256']} – {e['url_origem'] or e['nome_original']}", style="List Bullet")
    d.add_heading("Metodologia", 1)
    for m in data["metodologia"]:
        d.add_paragraph(m, style="List Bullet")
    d.save(str(path))


def render_pdf(data: dict, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape
    styles = getSampleStyleSheet()
    flow = []
    P = lambda t, s="BodyText": flow.append(Paragraph(escape(str(t)), styles[s]))  # noqa: E731
    p = data["projeto"]
    P(f"{data['titulo']}: {p['name']}", "Title")
    P(f"Gerado em {data['gerado_em']} · Finalidade: {p['purpose']}", "Italic")
    P(data["aviso"], "Italic")
    flow.append(Spacer(1, .4 * cm))
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    rows: list[list] = []

    def flush():
        if rows:
            cell = styles["BodyText"].clone("cel", fontSize=7.5, leading=9)
            t = Table([[Paragraph(escape(c), cell) for c in r] for r in rows], repeatRows=1,
                      colWidths=[1 * cm, 8.5 * cm, 4.5 * cm, 1.3 * cm, 1.5 * cm][:len(rows[0])])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .3, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f8"))]))
            flow.append(t)
            rows.clear()

    for line in render_markdown(data).splitlines()[6:]:
        if line.startswith("|"):
            if not line.startswith("|---"):
                rows.append([c.strip() for c in line.strip("|").split("|")])
            continue
        flush()
        if line.startswith("## "):
            P(line[3:], "Heading2")
        elif line.startswith("### "):
            P(line[4:], "Heading3")
        elif line.strip():
            P(line.lstrip("- ").replace("**", "").replace("`", ""))
    flush()
    SimpleDocTemplate(str(path), pagesize=A4, title=f"{data['titulo']} {p['name']}",
                      leftMargin=1.6 * cm, rightMargin=1.6 * cm).build(flow)


def generate(project_id: int, fmt: str = "html", tipo: str = "completo") -> Path:
    data = collect(project_id, tipo)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir = get_settings().reports_dir / f"projeto-{project_id}"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{tipo}-{stamp}.{fmt}"
    if fmt == "html":
        path.write_text(render_html(data), encoding="utf-8")
    elif fmt == "json":
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    elif fmt == "md":
        path.write_text(render_markdown(data), encoding="utf-8")
    elif fmt == "docx":
        render_docx(data, path)
    elif fmt == "pdf":
        render_pdf(data, path)
    else:
        raise ValueError(f"Formato não suportado: {fmt}")
    from . import cofre
    if cofre.is_active():
        path.write_bytes(cofre.encrypt_bytes(path.read_bytes()))
    audit.log("relatorio_gerado", project_id=project_id, target=path.name, details={"formato": fmt, "tipo": tipo})
    return path


def read_file(path: Path) -> bytes:
    """Conteúdo de um relatório gravado, decifrado se o cofre estiver ativo."""
    from . import cofre
    return cofre.decrypt_bytes(path.read_bytes())
