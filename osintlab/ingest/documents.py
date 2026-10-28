"""Extração de texto e metadados de documentos (PDF, DOCX, XLSX, CSV, JSON, TXT, HTML, imagens)."""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Extracted:
    text: str
    title: str = ""
    metadata: dict = field(default_factory=dict)
    kind: str = "documento"
    warnings: list[str] = field(default_factory=list)


def detect_type(name: str, mime: str | None, data: bytes) -> str:
    ext = Path(name or "").suffix.lower()
    if data[:5] == b"%PDF-" or ext == ".pdf":
        return "pdf"
    if data[:2] == b"PK":
        if ext == ".xlsx" or b"xl/" in data[:4000]:
            return "xlsx"
        if ext == ".docx" or b"word/" in data[:4000]:
            return "docx"
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n" or ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".heic", ".gif"):
        return "imagem"
    if ext in (".mp3", ".wav", ".m4a", ".ogg", ".flac", ".mp4", ".mkv", ".mov", ".webm"):
        return "audiovisual"
    if ext == ".csv" or (mime or "").endswith("csv"):
        return "csv"
    if ext == ".json" or (mime or "").endswith("json"):
        return "json"
    head = data[:2000].lstrip().lower()
    if ext in (".xml", ".rss", ".atom") or head.startswith(b"<?xml") and (b"<rss" in head or b"<feed" in head):
        return "feed" if (b"<rss" in head or b"<feed" in head) else "xml"
    if ext in (".html", ".htm") or head.startswith(b"<!doctype html") or b"<html" in head:
        return "html"
    return "texto"


def extract_pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(f"[pág. {i + 1}]\n" + (page.extract_text() or ""))
        except Exception as e:  # páginas corrompidas não interrompem o lote
            pages.append(f"[pág. {i + 1}] (erro: {e})")
    meta = {}
    if reader.metadata:
        for k, v in reader.metadata.items():
            meta[str(k).lstrip("/")] = str(v)
    meta["paginas"] = len(reader.pages)
    text = "\n".join(pages)
    warnings = []
    if len(text.strip()) < 30 * max(1, len(reader.pages)):
        ocr_text, info = ocr_pdf(data)
        if ocr_text.strip():
            text = ocr_text
            meta["ocr"] = info
            warnings.append(f"Texto obtido por OCR ({info['paginas']} pág., {info['motor']}); confira trechos críticos "
                            "no original.")
        else:
            warnings.append("Pouco texto extraído: o PDF parece digitalizado. " + info.get("nota", ""))
    return Extracted(text=text, title=meta.get("Title", ""), metadata=meta, warnings=warnings)


def ocr_pdf(data: bytes, lang: str = "por+eng") -> tuple[str, dict]:
    """OCR página a página: renderiza com pypdfium2 (300 dpi) e reconhece com o motor ativo (Windows ou Tesseract)."""
    import os

    from . import ocr
    max_pages = int(os.environ.get("OSINTLAB_OCR_MAX_PAGINAS", "50"))
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return "", {"nota": "Para OCR de PDF instale o extra: pip install -e \".[ocr]\"."}
    if not ocr.disponivel():
        st = ocr.estado()
        return "", {"nota": "OCR desligado em Configurações." if st["preferencia"] == "nenhum" else
                    "Nenhum motor de OCR disponível: no Windows 10/11 o OCR do sistema é usado automaticamente; "
                    "em outros sistemas instale o Tesseract com o idioma 'por'."}
    pdf = pdfium.PdfDocument(data)
    out, motor = [], ""
    n = min(len(pdf), max_pages)
    for i in range(n):
        img = pdf[i].render(scale=300 / 72).to_pil()
        texto, motor = ocr.reconhecer(img)
        out.append(f"[pág. {i + 1} — OCR]\n" + texto)
    info = {"paginas": n, "motor": motor}
    if len(pdf) > n:
        info["truncado"] = f"{len(pdf) - n} página(s) além do limite OSINTLAB_OCR_MAX_PAGINAS"
    return "\n".join(out), info


def extract_docx(data: bytes) -> Extracted:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    cp = d.core_properties
    meta = {k: str(getattr(cp, k)) for k in ("author", "created", "modified", "last_modified_by", "title",
                                              "subject", "keywords", "revision") if getattr(cp, k, None)}
    return Extracted(text="\n".join(parts), title=cp.title or "", metadata=meta)


def extract_xlsx(data: bytes, max_rows: int = 20000) -> Extracted:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, meta = [], {"planilhas": wb.sheetnames}
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append("... (truncado)")
                break
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    props = wb.properties
    for k in ("creator", "created", "modified", "lastModifiedBy", "title"):
        v = getattr(props, k, None)
        if v:
            meta[k] = str(v)
    return Extracted(text="\n".join(parts), title=meta.get("title", ""), metadata=meta, kind="dataset")


def extract_csv(data: bytes) -> Extracted:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    meta = {"linhas": len(rows), "colunas": rows[0] if rows else []}
    return Extracted(text="\n".join(" | ".join(r) for r in rows), metadata=meta, kind="dataset")


def extract_json(data: bytes) -> Extracted:
    obj = json.loads(data.decode("utf-8-sig", errors="replace"))
    flat: list[str] = []

    def walk(o, path=""):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, f"{path}.{k}" if path else str(k))
        elif isinstance(o, list):
            for i, v in enumerate(o[:5000]):
                walk(v, f"{path}[{i}]")
        else:
            flat.append(f"{path}: {o}")

    walk(obj)
    return Extracted(text="\n".join(flat), metadata={"tipo_raiz": type(obj).__name__}, kind="dataset")


def extract_image(data: bytes) -> Extracted:
    from .images import analyze_image
    info = analyze_image(data)
    text = info.pop("ocr_text", "") or ""
    return Extracted(text=text, metadata=info, kind="imagem",
                     warnings=[] if text else ["Sem texto (OCR indisponível ou imagem sem texto)."])


def extract(data: bytes, name: str = "", mime: str | None = None) -> Extracted:
    t = detect_type(name, mime, data)
    if t == "pdf":
        return extract_pdf(data)
    if t == "docx":
        return extract_docx(data)
    if t == "xlsx":
        return extract_xlsx(data)
    if t == "csv":
        return extract_csv(data)
    if t == "json":
        return extract_json(data)
    if t == "imagem":
        return extract_image(data)
    if t == "html":
        from .web import extract_html
        return extract_html(data.decode("utf-8", errors="replace"))
    if t == "feed":
        from .rss import parse_feed
        entries = parse_feed(data)
        return Extracted(text="\n\n".join(f"{e['title']}\n{e['link']}\n{e['summary']}" for e in entries),
                         metadata={"entradas": len(entries)}, kind="rss")
    if t == "audiovisual":
        from ..ai.transcribe import transcribe
        res = transcribe(data, name)
        return Extracted(text=res.get("text", ""), metadata=res.get("metadata", {}), kind="audiovisual",
                         warnings=res.get("warnings", []))
    return Extracted(text=data.decode("utf-8", errors="replace"), kind="documento")
