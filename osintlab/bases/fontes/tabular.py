"""Importador genérico de bases tabulares (CSV, CSV dentro de ZIP, XLSX) com versionamento.

Versão remota = ETag/Last-Modified do servidor (sem baixar) ou, na falta deles, o SHA-256
do conteúdo. Diferenças entre versões completas são calculadas por chave + hash da linha.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from ..catalog import Fonte, RemoteInfo, SearchHit
from ..download import fetch_file, head
from ..store import set_meta
from ._util import fold, fts_query

csv.field_size_limit(10 ** 8)


def _col(name: str) -> str:
    c = re.sub(r"\W+", "_", fold(name)).strip("_")
    return (c or "col")[:60]


def iter_rows(path: Path, *, sep: str | None = None, encoding: str = "utf-8-sig", member_rx: str | None = None,
              header: bool = True, columns: list[str] | None = None, skip_rows: int = 0):
    """Gera (cabeçalho, linhas) de CSV, ZIP com CSVs ou XLSX."""
    name = path.name.lower()
    if name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        for _ in range(skip_rows):
            next(it, None)
        hdr = columns or [str(c or f"col{i}") for i, c in enumerate(next(it))]
        yield hdr, ([("" if v is None else str(v)) for v in r] for r in it)
        return
    streams = []
    if zipfile.is_zipfile(path):
        z = zipfile.ZipFile(path)
        members = [m for m in z.namelist() if not m.endswith("/") and (not member_rx or re.search(member_rx, m, re.I))]
        streams = [(m, z.open(m)) for m in members]
    else:
        streams = [(path.name, path.open("rb"))]
    for _, raw in streams:
        text = io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline="")
        sample = text.read(20000)
        delim = sep
        if not delim:
            try:
                delim = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
            except csv.Error:
                delim = ";"
        chained = _Chain(sample, text)
        reader = csv.reader(chained, delimiter=delim)
        for _ in range(skip_rows):
            next(reader, None)
        hdr = columns or (next(reader) if header else None)
        yield hdr, reader


class _Chain(io.TextIOBase):
    """Reinjeta a amostra lida para detecção do separador antes do restante do fluxo."""

    def __init__(self, first: str, rest):
        self._buf, self._rest = first, rest

    def readable(self):
        return True

    def read(self, n=-1):
        if self._buf:
            out, self._buf = self._buf, ""
            return out + (self._rest.read() if n == -1 else "")
        return self._rest.read(n)

    def __iter__(self):
        if self._buf:
            first, self._buf = self._buf, ""
            lines = first.splitlines(keepends=True)
            if lines and not lines[-1].endswith(("\n", "\r")):
                lines[-1] += self._rest.readline()
            yield from lines
        yield from self._rest


def import_table(conn, path: Path, table: str, *, key: str | None = None, text_cols: list[str] | None = None,
                 drop: list[str] | None = None, **parse) -> dict:
    """Importa para `table` com colunas normalizadas + `_hash` + índice FTS do texto.
    `drop`: colunas descartadas na importação (minimização de dados pessoais)."""
    n = 0
    created = False
    dropset = {_col(c) for c in (drop or [])}
    for hdr, rows in iter_rows(path, **parse):
        keep = [i for i, h in enumerate(hdr) if _col(h) not in dropset]
        hdr = [hdr[i] for i in keep]
        rows = ([r[i] if i < len(r) else "" for i in keep] for r in rows)
        cols = []
        for h in hdr:
            c = _col(h)
            while c in cols:
                c += "_"
            cols.append(c)
        if not created:
            conn.execute(f'CREATE TABLE IF NOT EXISTS "{table}" (_linha INTEGER PRIMARY KEY, _chave TEXT, _hash TEXT, '
                         + ", ".join(f'"{c}" TEXT' for c in cols) + ")")
            conn.execute(f'CREATE VIRTUAL TABLE IF NOT EXISTS "{table}_fts" USING fts5(texto, content="", '
                         "tokenize='unicode61 remove_diacritics 2')")
            created = True
        kidx = cols.index(_col(key)) if key and _col(key) in cols else None
        tidx = [cols.index(_col(c)) for c in (text_cols or []) if _col(c) in cols] or list(range(len(cols)))
        batch, fts = [], []
        q = f'INSERT INTO "{table}" (_chave, _hash, ' + ", ".join(f'"{c}"' for c in cols) + ") VALUES (" + \
            ",".join("?" * (len(cols) + 2)) + ")"
        for r in rows:
            r = (list(r) + [""] * len(cols))[:len(cols)]
            h = hashlib.sha1("\x1f".join(r).encode()).hexdigest()
            batch.append([r[kidx] if kidx is not None else None, h, *r])
            if len(batch) >= 5000:
                conn.executemany(q, batch)
                n += len(batch)
                batch.clear()
        if batch:
            conn.executemany(q, batch)
            n += len(batch)
        # índice textual (contentless) ligado ao rowid
        conn.execute(f'INSERT INTO "{table}_fts"(rowid, texto) SELECT _linha, ' +
                     " || ' ' || ".join(f'coalesce("{cols[i]}","")' for i in tidx) + f' FROM "{table}"')
    if key:
        conn.execute(f'CREATE INDEX IF NOT EXISTS "ix_{table}_chave" ON "{table}"(_chave)')
    return {"linhas": n}


def diff_tables(conn, old_path: Path | None, table: str) -> dict:
    if not old_path or not old_path.exists():
        return {"added": conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0], "modified": 0, "deleted": 0}
    from ..store import ro_uri
    conn.execute("ATTACH DATABASE ? AS antiga", (ro_uri(old_path),))
    try:
        has_key = conn.execute(f'SELECT COUNT(_chave) FROM "{table}"').fetchone()[0] > 0
        if has_key:
            added = conn.execute(f'SELECT COUNT(*) FROM "{table}" n WHERE NOT EXISTS (SELECT 1 FROM antiga."{table}" o '
                                 "WHERE o._chave=n._chave)").fetchone()[0]
            deleted = conn.execute(f'SELECT COUNT(*) FROM antiga."{table}" o WHERE NOT EXISTS (SELECT 1 FROM "{table}" n '
                                   "WHERE n._chave=o._chave)").fetchone()[0]
            modified = conn.execute(f'SELECT COUNT(*) FROM "{table}" n JOIN antiga."{table}" o ON o._chave=n._chave '
                                    "WHERE o._hash != n._hash").fetchone()[0]
        else:
            added = conn.execute(f'SELECT COUNT(*) FROM "{table}" n WHERE NOT EXISTS (SELECT 1 FROM antiga."{table}" o '
                                 "WHERE o._hash=n._hash)").fetchone()[0]
            deleted = conn.execute(f'SELECT COUNT(*) FROM antiga."{table}" o WHERE NOT EXISTS (SELECT 1 FROM "{table}" n '
                                   "WHERE n._hash=o._hash)").fetchone()[0]
            modified = 0
    except Exception:
        added, modified, deleted = conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0], 0, 0
    finally:
        conn.commit()
        conn.execute("DETACH DATABASE antiga")
    return {"added": added, "modified": modified, "deleted": deleted}


class TabularFonte(Fonte):
    """Fonte B baseada em um ou mais arquivos tabulares publicados em URLs estáveis."""

    arquivos: list[dict] = []   # [{"url":..., "tabela":..., "chave":..., "texto": [...], "parse": {...}}]
    titulo_cols: list[str] = []
    sincronizavel = True
    entradas = ["texto"]

    def urls(self, ctx) -> list[dict]:
        return self.arquivos

    def _dest(self, ctx, a: dict) -> Path:
        return ctx.workdir() / (a.get("nome") or Path(a["url"].split("?")[0]).name or "dados")

    def remote_info(self, ctx) -> RemoteInfo:
        """Usa ETag/Last-Modified quando o servidor os fornece; senão baixa e usa o SHA-256 do conteúdo."""
        parts, total = [], 0
        for a in self.urls(ctx):
            h = head(a["url"])
            if h["status"] >= 400:
                raise RuntimeError(f"HTTP {h['status']} em {a['url']}")
            marker = h.get("etag") or h.get("last_modified")
            if not marker:
                d = fetch_file(a["url"], self._dest(ctx, a), cancel=ctx.cancel)
                ctx.add_file(d.url, d.sha256, d.size, d.path.name)
                marker = "sha256:" + d.sha256
            parts.append(marker)
            total += h.get("size") or 0
        version = hashlib.sha1("|".join(parts).encode()).hexdigest()[:16]
        return RemoteInfo(version=version, size=total or None, extra={"marcadores": parts})

    def full_import(self, ctx, conn, remote: RemoteInfo) -> dict:
        from ...db import get_db
        from ..store import version_path
        st = get_db().one("SELECT current_version FROM bases_state WHERE source_id=?", (self.id,))
        prev = version_path(self.id, st["current_version"]) if st and st["current_version"] else None
        totals = {"added": 0, "modified": 0, "deleted": 0}
        for a in self.urls(ctx):
            dest = self._dest(ctx, a)
            if not dest.exists():  # pode já ter sido baixado na verificação de versão
                d = fetch_file(a["url"], dest, cancel=ctx.cancel, expected_size=remote.size)
                ctx.add_file(d.url, d.sha256, d.size, dest.name)
            import_table(conn, dest, a["tabela"], key=a.get("chave"), text_cols=a.get("texto"), drop=a.get("descartar"),
                         **a.get("parse", {}))
            dd = diff_tables(conn, prev, a["tabela"])
            for k in totals:
                totals[k] += dd[k]
            dest.unlink(missing_ok=True)
        set_meta(conn, tabela_principal=self.urls(ctx)[0]["tabela"], schema_versao=f"{self.id}-tab-1",
                 arquivos=json.dumps([a["url"] for a in self.urls(ctx)]))
        return {"mensagem": f"{len(self.urls(ctx))} arquivo(s) importado(s)", **totals}

    def validate(self, conn) -> dict:
        out = {"ok": True, "problemas": []}
        for a in self.arquivos:
            n = conn.execute(f'SELECT COUNT(*) FROM "{a["tabela"]}"').fetchone()[0]
            out[a["tabela"]] = n
            if n == 0:
                out["ok"] = False
                out["problemas"].append(f"tabela {a['tabela']} vazia")
        return out

    def hit_from_row(self, table: str, row) -> SearchHit:
        d = {k: row[k] for k in row.keys() if not k.startswith("_")}
        title = " · ".join(str(d[c]) for c in self.titulo_cols if d.get(c)) or str(row["_chave"] or row["_linha"])
        return SearchHit(fonte=self.id, registro_id=f"{table}:{row['_linha']}", titulo=title[:200], tipo=table,
                         resumo="; ".join(f"{k}: {v}" for k, v in list(d.items())[:8] if v)[:400], url=self.url_oficial,
                         identificadores=[str(row["_chave"])] if row["_chave"] else [], nomes=[title], atributos=d)

    def search_local(self, conn, query, kind, limit=25):
        from ..store import get_meta
        version = get_meta(conn).get("versao")
        out = []
        for a in self.arquivos:
            t = a["tabela"]
            if a.get("chave") and kind in ("identificador", "cnpj", "codigo"):
                key = re.sub(r"\D", "", query) if kind == "cnpj" else query
                rows = conn.execute(f'SELECT * FROM "{t}" WHERE replace(replace(replace(_chave,".",""),"/",""),"-","")=? '
                                    f"OR _chave=? LIMIT ?", (key, query, limit)).fetchall()
            else:
                rows = conn.execute(f'SELECT t.* FROM "{t}_fts" f JOIN "{t}" t ON t._linha=f.rowid '
                                    f'WHERE "{t}_fts" MATCH ? LIMIT ?', (fts_query(query), limit)).fetchall()
            for r in rows:
                h = self.hit_from_row(t, r)
                h.versao_base = version
                out.append(h)
        return out[:limit]
