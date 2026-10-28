"""Dossiê com um clique: pesquisa unificada + rede societária + checagem das pessoas e empresas ligadas nas listas
de sanções, fiscalização, PEP, procurados e justiça, num documento único (tela, HTML, Markdown) que pode ser salvo no projeto.

Os resultados mantêm os níveis da pesquisa (documental, provável, hipótese…): um nome igual numa lista nunca é
tratado como a mesma pessoa. Ausência de resultado não comprova inexistência.
"""
from __future__ import annotations

import secrets
import threading
import time
from concurrent.futures import ThreadPoolExecutor

from . import audit, pesquisa, vinculos
from .bases.catalog import TEMAS, fontes

TEMAS_ALERTA = ("sancoes", "procurados_alertas", "justica_seguranca", "fiscalizacao", "pep")
ORDEM_TEMAS = ["empresas", "sancoes", "fiscalizacao", "procurados_alertas", "justica_seguranca", "pep", "transparencia", "processos",
               "eleitoral", "saude", "beneficios", "veiculos", "aeronaves", "geografia", "infraestrutura", "noticias"]
MAX_LIGADOS = 15
_CACHE: dict[str, tuple[float, dict]] = {}
_LOCK = threading.Lock()
VALIDADE_S = 3600


def _secoes(resultados: list[dict]) -> list[dict]:
    por_tema: dict[str, list[dict]] = {}
    for h in resultados:
        f = fontes().get(h["fonte"])
        por_tema.setdefault(f.tema if f else "outros", []).append(h)
    ordem = ORDEM_TEMAS + sorted(t for t in por_tema if t not in ORDEM_TEMAS)
    return [{"tema": t, "titulo": TEMAS.get(t, "Outras fontes"), "resultados": por_tema[t]} for t in ordem if t in por_tema]


def _ligados(rede: dict) -> list[dict]:
    """Pessoas e empresas da rede (exceto a raiz), priorizando sócios diretos e quem aparece em mais empresas."""
    raiz = rede.get("raiz")
    diretos = {e["de"] for e in rede["arestas"] if e["para"] == raiz}
    nos = [n for n in rede["nos"] if n["id"] != raiz]
    nos.sort(key=lambda n: (n["id"] not in diretos, -n.get("empresas", 0), n["rotulo"]))
    return [{"nome": n["rotulo"], "tipo": n["tipo"], "cnpj": n.get("cnpj", "")} for n in nos[:MAX_LIGADOS]]


def _checar(ligado: dict, selecionadas: list[str], pid: int | None) -> list[dict]:
    consulta = ligado["cnpj"] or ligado["nome"]
    forcar = "" if ligado["cnpj"] else ("pessoa" if ligado["tipo"] == "pessoa" else "organizacao")
    r = pesquisa.run(consulta, project_id=pid, selected=selecionadas, confirmado=True, forcar=forcar, limit=5)
    return [h | {"ligado": ligado["nome"]} for h in r["resultados"]
            if h["nivel"] in ("documental", "corroborada", "provavel") and fontes()[h["fonte"]].tema in TEMAS_ALERTA]


def gerar(consulta: str, *, project_id: int | None = None, confirmado: bool = False, finalidade: str = "",
          profundidade: int = 2, campo: str = "") -> dict:
    inicio = time.time()
    res = pesquisa.run(consulta, project_id=project_id, confirmado=confirmado, forcar=campo)
    d = {"consulta": consulta, "gerado_em": time.strftime("%Y-%m-%d %H:%M"), "finalidade": finalidade,
         "pesquisa": res, "requer_confirmacao": res.get("requer_confirmacao", False), "rede": None, "alertas": [],
         "ligados": [], "erros": list(res.get("erros", []))}
    if d["requer_confirmacao"] or {"cpf", "invalido"} & set(res["consulta"]["tipos"]):
        return d
    d["secoes"] = _secoes(res["resultados"])
    if "cnpj" in res["consulta"]["tipos"]:
        try:
            d["rede"] = vinculos.rede(res["consulta"]["normalizado"], profundidade)
        except Exception as e:  # noqa: BLE001 — o dossiê continua sem a rede
            d["erros"].append({"fonte": "rede_societaria", "erro": f"{e.__class__.__name__}: {e}"})
    if d["rede"] and d["rede"].get("encontrada"):
        d["ligados"] = _ligados(d["rede"])
        sel = [f.id for f in fontes().values() if f.tema in TEMAS_ALERTA]
        with ThreadPoolExecutor(max_workers=5) as ex:
            for hits in ex.map(lambda x: _checar(x, sel, project_id), d["ligados"]):
                d["alertas"].extend(hits)
    from . import riscos
    d["riscos"] = riscos.avaliar(d)
    d["resumo"] = _resumo(d)
    d["duracao_s"] = round(time.time() - inicio, 1)
    audit.log("dossie_gerado", project_id=project_id,
              details={"tipos": res["consulta"]["tipos"], "resultados": len(res["resultados"]),
                       "ligados": len(d["ligados"]), "alertas": len(d["alertas"])})
    return d


def de_pesquisa(res: dict, consulta: str, finalidade: str = "") -> dict:
    """Dossiê simples a partir de uma pesquisa já feita (sem refazer consultas): exportação direto da tela."""
    d = {"consulta": consulta, "gerado_em": time.strftime("%Y-%m-%d %H:%M"), "finalidade": finalidade, "pesquisa": res,
         "requer_confirmacao": False, "rede": None, "alertas": [], "ligados": [], "erros": list(res.get("erros", [])),
         "secoes": _secoes(res.get("resultados", []))}
    from . import riscos
    d["riscos"] = riscos.avaliar(d)
    d["resumo"] = _resumo(d)
    return d


def _resumo(d: dict) -> list[str]:
    res = d["pesquisa"]
    linhas = [res.get("sintese") or f"{len(res['resultados'])} registro(s) nas fontes consultadas."]
    if res.get("nomes_cpf"):
        linhas.append("Nomes encontrados com o CPF parcial (conferidos por nome e CPF): " + "; ".join(res["nomes_cpf"]) + ".")
    if r := d["rede"]:
        if r.get("encontrada"):
            emp = sum(1 for n in r["nos"] if n["tipo"] == "empresa") - 1
            soc = sum(1 for e in r["arestas"] if e["para"] == r["raiz"])
            linhas.append(f"Rede societária ({r['origem']}): {soc} sócio(s) direto(s), {emp} empresa(s) ligada(s)"
                          + (f"; {len(r['comuns'])} pessoa(s) em mais de uma empresa" if r["comuns"] else "") + ".")
        elif r.get("aviso"):
            linhas.append(r["aviso"])
    if d.get("riscos"):
        n = {g: sum(1 for i in d["riscos"] if i["gravidade"] == g) for g in ("alta", "media", "baixa")}
        linhas.append(f"{len(d['riscos'])} indicador(es) de risco: {n['alta']} de gravidade alta, {n['media']} média, "
                      f"{n['baixa']} baixa.")
    if d["ligados"]:
        linhas.append(f"{len(d['ligados'])} pessoa(s)/empresa(s) ligada(s) conferida(s) em sanções, fiscalização, PEP, "
                      f"procurados e justiça: {len(d['alertas'])} correspondência(s) que merecem verificação.")
    return [x for x in linhas if x]


# ------------------------------------------------------------------ guarda temporária (exportação e salvamento)
def guardar(d: dict) -> str:
    token = secrets.token_urlsafe(12)
    with _LOCK:
        agora = time.time()
        for k in [k for k, (t, _) in _CACHE.items() if agora - t > VALIDADE_S]:
            _CACHE.pop(k, None)
        _CACHE[token] = (agora, d)
    return token


def obter(token: str) -> dict | None:
    with _LOCK:
        v = _CACHE.get(token)
    return v[1] if v and time.time() - v[0] <= VALIDADE_S else None


# ------------------------------------------------------------------ Markdown
def markdown(d: dict, projeto: dict | None = None) -> str:
    niveis = pesquisa.NIVEIS
    out = [f"# Dossiê: {d['consulta']}", "", f"Gerado em {d['gerado_em']} pelo OSINT Lab."]
    if projeto:
        out.append(f"Caso: {projeto['name']}. Finalidade: {projeto['purpose']}.")
    if d.get("finalidade"):
        out.append(f"Finalidade declarada desta consulta: {d['finalidade']}.")
    out += ["", "## Resumo", ""] + [f"- {x}" for x in d.get("resumo", [])]
    if d.get("riscos"):
        out += ["", "## Indicadores de risco (sinais para conferir, não conclusões)", ""]
        for i in d["riscos"]:
            out.append(f"- **[{i['gravidade'].upper()}] {i['titulo']}**: {i['explicacao']} "
                       + ("Evidências: " + "; ".join(i["evidencias"]) if i["evidencias"] else ""))
    for s in d.get("secoes", []):
        out += ["", f"## {s['titulo']}", ""]
        for h in s["resultados"]:
            out.append(f"- **{h['titulo']}** ({fontes()[h['fonte']].nome}; {niveis.get(h['nivel'], h['nivel'])}"
                       + (f"; via nome {h['via_nome']}" if h.get("via_nome") else "") + f"): {h.get('resumo', '')}"
                       + (f" <{h['url']}>" if h.get("url") else ""))
    if (r := d.get("rede")) and r.get("encontrada"):
        nos = {n["id"]: n for n in r["nos"]}
        out += ["", "## Rede societária", "", r["aviso"], ""]
        for e in r["arestas"]:
            a, b = nos[e["de"]], nos[e["para"]]
            out.append(f"- {a['rotulo']} → {b['rotulo']} ({b.get('cnpj', '')}): {e['qualificacao']}, desde "
                       f"{e['desde'] or '?'} [{e['nivel']}]")
        if r["comuns"]:
            out += ["", "Pessoas em mais de uma empresa:"] + [f"- {c['nome']} ({c['documento']}): {c['empresas']} empresas"
                                                             for c in r["comuns"]]
    if d.get("ligados"):
        out += ["", "## Pessoas e empresas ligadas: sanções, fiscalização, PEP, procurados e justiça", ""]
        if d["alertas"]:
            out += [f"- {h['ligado']}: **{h['titulo']}** ({fontes()[h['fonte']].nome}; "
                    f"{niveis.get(h['nivel'], h['nivel'])}). {h.get('justificativa', '')}" for h in d["alertas"]]
        else:
            out.append("Nenhuma correspondência nas fontes consultadas.")
        out.append("Conferidos: " + "; ".join(x["nome"] for x in d["ligados"]) + ".")
    if comp := [c for c in d["pesquisa"].get("complementares") or [] if not c.get("erro")]:
        out += ["", "## Fontes online complementares", ""]
        for c in comp:
            out += [f"### {c['nome']}", "", (c.get("texto") or "")[:1500], ""]
    if assist := d["pesquisa"].get("assistidas"):
        out += ["", "## Consultas assistidas sugeridas", ""] + [f"- {a['nome']}: <{a['url']}>" for a in assist]
    if d.get("erros"):
        out += ["", "## Fontes com falha", ""] + [f"- {e['fonte']}: {e['erro']}" for e in d["erros"]]
    out += ["", "---", "Vínculos e correspondências são registros declaratórios das fontes na data indicada e não "
            "implicam responsabilidade. Nome igual não é identidade. Ausência de resultado não comprova inexistência."]
    return "\n".join(out) + "\n"


def salvar_no_projeto(pid: int, d: dict, selecionados: list[dict] | None = None) -> dict:
    """Guarda o dossiê no caso. Os registros salvos um a um são os escolhidos ou, sem escolha, os relevantes."""
    from .core import add_item
    relevantes = selecionados if selecionados is not None else [
        h for h in d["pesquisa"]["resultados"] + d.get("alertas", [])
        if h["nivel"] not in ("hipotese", "contraditorio")]  # o texto do dossiê guarda todos
    ids = pesquisa.save_hits(pid, relevantes, d["consulta"])
    rede_item = None
    if (r := d.get("rede")) and r.get("encontrada"):
        rede_item = vinculos.salvar_no_projeto(pid, r)
    iid = add_item(pid, kind="dossie", title=f"Dossiê: {d['consulta']} ({d['gerado_em']})",
                   text=markdown(d), metadata={"dossie": {"consulta": d["consulta"], "gerado_em": d["gerado_em"],
                                                          "registros": ids, "rede": rede_item}},
                   collector="dossie")
    return {"item": iid, "registros": len(ids), "rede": rede_item}


def _linhas_md(md: str):
    """(tipo, texto) de um Markdown simples: título, seção, subseção, item, parágrafo."""
    for linha in md.splitlines():
        t = linha.rstrip()
        if not t or t == "---":
            continue
        if t.startswith("# "):
            yield "titulo", t[2:]
        elif t.startswith("## "):
            yield "secao", t[3:]
        elif t.startswith("### "):
            yield "subsecao", t[4:]
        elif t.startswith("- "):
            yield "item", t[2:]
        else:
            yield "texto", t


def _limpo(t: str) -> str:
    import re
    return re.sub(r"<(https?://[^>]+)>", r"\1", t.replace("**", "").replace("`", ""))


def _fontes_pdf() -> tuple[str, str]:
    """Fonte Unicode (Vera, distribuída com o reportlab): acentos, travessões e setas saem corretos."""
    import os

    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pasta = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    try:
        if "Vera" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("Vera", os.path.join(pasta, "Vera.ttf")))
            pdfmetrics.registerFont(TTFont("VeraBd", os.path.join(pasta, "VeraBd.ttf")))
        return "Vera", "VeraBd"
    except Exception:  # sem os arquivos: fontes padrão (acentos funcionam; travessões não)
        return "Helvetica", "Helvetica-Bold"


def pdf(d: dict, projeto: dict | None = None) -> bytes:
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    st = getSampleStyleSheet()
    fonte, negrito = _fontes_pdf()
    for nome in ("Title", "Heading2", "Heading3", "BodyText"):
        st[nome].fontName = negrito if nome != "BodyText" else fonte
    estilo = {"titulo": "Title", "secao": "Heading2", "subsecao": "Heading3", "item": "BodyText", "texto": "BodyText"}
    flow = []
    for tipo, t in _linhas_md(markdown(d, projeto)):
        t = escape(_limpo(t))
        flow.append(Paragraph(("• " if tipo == "item" else "") + t, st[estilo[tipo]]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=f"Dossiê {d['consulta']}", leftMargin=1.6 * cm,
                      rightMargin=1.6 * cm).build(flow)
    return buf.getvalue()


def docx(d: dict, projeto: dict | None = None) -> bytes:
    import io

    import docx as _docx
    doc = _docx.Document()
    for tipo, t in _linhas_md(markdown(d, projeto)):
        t = _limpo(t)
        if tipo == "titulo":
            doc.add_heading(t, 0)
        elif tipo == "secao":
            doc.add_heading(t, 1)
        elif tipo == "subsecao":
            doc.add_heading(t, 2)
        elif tipo == "item":
            doc.add_paragraph(t, style="List Bullet")
        else:
            doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
